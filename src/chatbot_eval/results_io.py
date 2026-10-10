from __future__ import annotations

import csv
import json
import logging
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .models import ChatbotResult, ExpectedBehavior, SilverQuestion, SourceRef
from .progress import track
from .response_errors import placeholder_error
from .validation import normalize_question_payload, parse_bool, validate_question_set
from .io import csv_question_row, question_sources

logger = logging.getLogger(__name__)


ALIASES = {
    "question": {"question", "prompt", "query", "user_prompt", "user prompt", "שאלה"},
    "expected_answer": {"expected_answer", "expected answer", "reference_answer", "reference answer", "gold_answer", "gold answer", "expected", "תשובה צפויה"},
    "answer": {"answer", "response", "chatbot_answer", "chatbot answer", "model_answer", "model answer", "bot_response", "bot response", "תשובה מהמודל"},
    "source": {"source", "sources", "evidence", "context", "retrieved_context", "retrieved context", "retrieved_chunks", "retrieved chunks", "מקור"},
    "id": {"id", "question_id", "question id", "external_id", "external id", "messageid", "message_id"},
    "topic": {"topic", "category", "נושא"},
    "answerable": {"answerable", "is_answerable", "is answerable"},
    "error": {"error", "exception", "failure"},
    "expected_behavior": {
        "expected_behavior", "expected behavior", "target_behavior", "target behavior",
    },
    "question_form": {"question_form", "question form", "form"},
    "parent_question_id": {
        "parent_question_id", "parent question id", "parent_id", "parent id",
        "source_question_id", "source question id",
    },
    "question_type": {"question_type", "question type", "type", "kind"},
    "difficulty": {"difficulty", "complexity"},
    "reference_claims": {
        "reference_claims", "reference claims", "gold_claims", "gold claims", "required_claims", "required claims",
    },
    "supporting_quotes": {
        "supporting_quotes", "supporting quotes", "evidence_quotes", "evidence quotes",
    },
    "reference_sources": {
        "reference_sources", "reference sources", "reference_sources_json", "reference sources json",
        "gold_sources", "gold sources", "gold_evidence", "gold evidence", "sources_json",
    },
}


@dataclass(frozen=True)
class ResultColumns:
    question: str | None = None
    expected_answer: str | None = None
    answer: str | None = None
    source: str | None = None
    id: str | None = None
    topic: str | None = None
    answerable: str | None = None
    error: str | None = None
    expected_behavior: str | None = None
    question_form: str | None = None
    parent_question_id: str | None = None
    question_type: str | None = None
    difficulty: str | None = None
    reference_claims: str | None = None
    supporting_quotes: str | None = None
    reference_sources: str | None = None


def _normalize(value: Any) -> str:
    return " ".join(str(value or "").strip().casefold().replace("_", " ").split())


def _resolve(headers: list[str], overrides: ResultColumns) -> dict[str, str | None]:
    blank_headers = [index + 1 for index, header in enumerate(headers) if not str(header).strip()]
    if blank_headers:
        raise ValueError(f"Blank column headers are not allowed (positions {blank_headers})")
    normalized_values = [_normalize(header) for header in headers]
    duplicate_headers = sorted(value for value, count in Counter(normalized_values).items() if count > 1)
    if duplicate_headers:
        raise ValueError(f"Duplicate column headers are not allowed: {duplicate_headers}")
    normalized = {_normalize(header): header for header in headers}
    resolved: dict[str, str | None] = {}
    for field, aliases in ALIASES.items():
        override = getattr(overrides, field, None)
        if override:
            if override not in headers:
                raise ValueError(f"Column {override!r} not found. Available columns: {headers}")
            resolved[field] = override
        else:
            matches = sorted({normalized[_normalize(alias)] for alias in aliases if _normalize(alias) in normalized})
            if len(matches) > 1:
                raise ValueError(
                    f"Multiple columns match {field!r}: {matches}. Use the explicit --{field.replace('_', '-')}-column option."
                )
            resolved[field] = matches[0] if matches else None
    missing = [field for field in ("question", "expected_answer", "answer") if not resolved[field]]
    if missing:
        raise ValueError(f"Could not identify required columns {missing}. Available columns: {headers}. Use explicit --*-column options.")
    return resolved


class _InputRow(dict):
    def __init__(self, values, source_row: int):
        super().__init__(values)
        self.source_row = source_row


def _validate_headers(headers: list[str]) -> None:
    if not headers or any(not str(header or "").strip() for header in headers):
        raise ValueError("Blank column headers are not allowed")
    normalized = [_normalize(header) for header in headers]
    if len(normalized) != len(set(normalized)):
        raise ValueError("Duplicate column headers are not allowed")


def _xlsx_rows(path: Path, sheet_name: str | None) -> tuple[str, list[dict[str, Any]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    if sheet_name and sheet_name not in workbook.sheetnames:
        raise ValueError(f"Sheet {sheet_name!r} not found. Available sheets: {workbook.sheetnames}")
    sheet = workbook[sheet_name] if sheet_name else workbook.active
    values = sheet.iter_rows(values_only=True)
    try:
        headers = [str(value).strip() if value is not None else "" for value in next(values)]
    except StopIteration as exc:
        raise ValueError(f"Workbook sheet {sheet.title!r} is empty") from exc
    _validate_headers(headers)
    rows = [_InputRow(zip(headers, row), number) for number, row in enumerate(values, 2) if any(value not in (None, "") for value in row)]
    workbook.close()
    return sheet.title, rows


def _tabular_rows(path: Path) -> tuple[str, list[dict[str, Any]]]:
    if path.suffix.lower() == ".jsonl":
        rows = []
        for line_number, line in enumerate(path.read_text(encoding="utf-8").splitlines(), 1):
            if not line.strip():
                continue
            try:
                value = json.loads(line)
                if not isinstance(value, dict):
                    raise ValueError(f"Every JSONL line must contain an object (line {line_number})")
                rows.append(_InputRow(value, line_number))
            except json.JSONDecodeError as exc:
                raise ValueError(f"Invalid JSONL at line {line_number}: {exc.msg}") from exc
        if not all(isinstance(row, dict) for row in rows):
            raise ValueError("Every JSONL line must contain an object")
        return "jsonl", rows
    with path.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        _validate_headers(reader.fieldnames or [])
        return "csv", [_InputRow(row, number) for number, row in enumerate(reader, 2)]


def _safe_error(value: str, category: str = "stored_error") -> str:
    compact = " ".join(value.split())
    if not compact:
        return ""
    return f"{category}: {compact[:240]}" + ("…" if len(compact) > 240 else "")


def _looks_like_error(answer: str, explicit_error: str) -> str:
    if explicit_error.strip():
        return _safe_error(explicit_error)
    if not answer.strip():
        return "Premade result contains an empty chatbot answer"
    if error := placeholder_error(answer):
        return error
    normalized = answer.casefold()
    if "spike arrest violation" in normalized:
        return "throttled: detected Spike Arrest violation in stored response"
    if normalized.lstrip().startswith(("{", "[")) and any(
        token in normalized for token in ("'fault'", '"fault"', "errorcode", "faultstring")
    ):
        return "api_fault: detected structured fault payload in stored response"
    return ""


def _as_bool(value: Any, default: bool = True) -> bool:
    return parse_bool(value, default)


@dataclass
class ImportDiagnostics:
    total: int = 0
    accepted: int = 0
    skipped: int = 0
    errors: list[dict[str, Any]] = field(default_factory=list)

    def as_dict(self) -> dict[str, Any]:
        return {"total": self.total, "accepted": self.accepted, "skipped": self.skipped, "errors": self.errors}


def read_premade_results(
    path: Path,
    *,
    sheet_name: str | None = None,
    columns: ResultColumns | None = None,
    progress_enabled: bool = False,
    strict: bool = False,
    diagnostics: ImportDiagnostics | None = None,
) -> list[tuple[SilverQuestion, ChatbotResult]]:
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        source_name, rows = _xlsx_rows(path, sheet_name)
    elif suffix in {".csv", ".jsonl"}:
        source_name, rows = _tabular_rows(path)
    else:
        raise ValueError("Premade results must be .xlsx, .csv, or .jsonl")
    if not rows:
        raise ValueError(f"No data rows found in {path}")
    headers = list(dict.fromkeys(key for row in rows for key in row if key is not None))
    mapping = _resolve(headers, columns or ResultColumns())
    logger.info("premade_results_loading path=%s source=%s row_count=%d columns=%r", path, source_name, len(rows), mapping)
    diagnostics = diagnostics if diagnostics is not None else ImportDiagnostics()
    diagnostics.total = len(rows)
    diagnostics.accepted = diagnostics.skipped = 0
    diagnostics.errors.clear()
    results = []
    external_ids: dict[str, str] = {}
    skipped: list[str] = []
    for index, row in enumerate(track(rows, enabled=progress_enabled, description="Reading premade results", total=len(rows)), 2):
        question_id = f"ROW{index:05d}"
        index = getattr(row, "source_row", index)
        try:
            if None in row and any(value not in (None, "") for value in (row.get(None) or [])):
                raise ValueError("row has more values than headers")
            if suffix == ".csv":
                row = csv_question_row(row)
            question_text = str(row.get(mapping["question"], "") or "").strip()
            expected = str(row.get(mapping["expected_answer"], "") or "").strip()
            answer = str(row.get(mapping["answer"], "") or "").strip()
            if not question_text or not expected:
                raise ValueError("missing question or reference answer")
            external_id = str(row.get(mapping["id"], "") or "").strip() if mapping["id"] else ""
            source = str(row.get(mapping["source"], "") or "").strip() if mapping["source"] else ""
            explicit_error = str(row.get(mapping["error"], "") or "").strip() if mapping["error"] else ""
            topic = str(row.get(mapping["topic"], "") or "לא סווג").strip() if mapping["topic"] else "לא סווג"
            payload = {key: row[key] for key in SilverQuestion.model_fields if key in row and key != "sources"}
            payload.update(id=question_id, topic=topic or "לא סווג", question=question_text,
                           expected_answer=expected, review_status=row.get("review_status") or "premade")
            if mapping["answerable"]:
                payload["answerable"] = row.get(mapping["answerable"])
            for field in (
                "expected_behavior", "question_form", "parent_question_id", "question_type", "difficulty",
                "reference_claims", "supporting_quotes",
            ):
                if mapping[field] and row.get(mapping[field]) is not None:
                    payload[field] = row.get(mapping[field])
            # `source` remains runtime retrieval. Gold evidence has explicit routes.
            if mapping["reference_sources"]:
                reference_column = mapping["reference_sources"]
                payload["sources"] = question_sources(row) if reference_column == "sources_json" else row.get(reference_column)
            elif any(key in row for key in ("sources_json", "source_files", "source_ids")):
                payload["sources"] = question_sources(row)
            question = SilverQuestion.model_validate(normalize_question_payload(payload))
            validate_question_set([question])
            if external_id and external_id in external_ids:
                raise ValueError("duplicate external question ID")
        except ValueError as exc:
            reason = f"row {index}: {exc}"
            skipped.append(reason)
            diagnostics.errors.append({"row": index, "reason": str(exc).split("\n", 1)[0]})
            logger.warning("premade_row_skipped row=%d reason=%s", index, exc)
            continue
        question_id = question.id
        if external_id:
            external_ids[external_id] = question_id
        if external_id and not question.reviewer_notes:
            question.reviewer_notes = f"Original ID: {external_id}"
        result = ChatbotResult(
            question_id=question_id, answer=answer, retrieved_context=source,
            metadata={"source_file": str(path), "source_sheet": source_name, "source_row": index, "original_id": external_id,
                      "reference_evidence_origin": "explicit_reference" if question.sources else "none"},
            error=_looks_like_error(answer, explicit_error),
        )
        results.append((question, result))
    diagnostics.accepted = len(results)
    diagnostics.skipped = len(skipped)
    for question, result in results:
        parent = question.parent_question_id
        if parent:
            result.metadata["original_parent_question_id"] = parent
            question.parent_question_id = external_ids.get(parent, "")
            if not question.parent_question_id:
                result.metadata["unresolved_parent_question_id"] = parent
    if strict and skipped:
        preview = "; ".join(skipped[:10])
        raise ValueError(
            f"Strict import rejected {len(skipped)} invalid row(s): {preview}"
            + (f"; and {len(skipped) - 10} more" if len(skipped) > 10 else "")
        )
    if not results:
        raise ValueError("No usable rows remain after import")
    logger.info("premade_results_loaded usable_rows=%d skipped_rows=%d", len(results), len(rows) - len(results))
    validate_question_set([question for question, _ in results])
    return results
