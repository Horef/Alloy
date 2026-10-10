from __future__ import annotations

import csv
import json
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentObject
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

from .progress import track


SUPPORTED_SUFFIXES = {".txt", ".md", ".rst", ".csv", ".json", ".jsonl", ".pdf", ".docx"}


def corpus_files(root: Path) -> list[Path]:
    """Supported documents under ``root``, skipping hidden files and folders (``.alloy/``, ``.git/``,
    editor settings), whose JSON or Markdown is tooling, not knowledge."""
    return [
        path for path in sorted(root.rglob("*"))
        if path.is_file() and path.suffix.lower() in SUPPORTED_SUFFIXES
        and not any(part.startswith(".") for part in path.relative_to(root).parts)
    ]


@dataclass(frozen=True)
class Chunk:
    id: str
    file: str
    location: str
    text: str


def _read_sections(path: Path) -> list[tuple[str, str]]:
    suffix = path.suffix.lower()
    if suffix == ".md":
        return _markdown_sections(path.read_text(encoding="utf-8", errors="replace"))
    if suffix in {".txt", ".rst"}:
        return [("document", path.read_text(encoding="utf-8", errors="replace"))]
    if suffix == ".csv":
        with path.open(encoding="utf-8-sig", errors="replace", newline="") as handle:
            return [("document", "\n\n".join(" | ".join(row) for row in csv.reader(handle)))]
    if suffix == ".json":
        return [("document", json.dumps(json.loads(path.read_text(encoding="utf-8")), ensure_ascii=False, indent=2))]
    if suffix == ".jsonl":
        return [("document", path.read_text(encoding="utf-8", errors="replace"))]
    if suffix == ".pdf":
        # The local PDF-to-Markdown converter restores Hebrew reading order and table cells. OCR is
        # off so ingestion is deterministic; convert scanned PDFs with `convert-pdf` first.
        from .pdf_markdown import convert_local

        return [(f"page {page.number}", page.markdown) for page in convert_local(path, ocr="off").pages]
    if suffix == ".docx":
        document = Document(path)
        # document.paragraphs and document.tables are separate collections and
        # therefore lose the meaningful paragraph/table/paragraph ordering.
        # Iterate the underlying XML body so a table is kept beside its text.
        ordered: list[str] = []
        for block in _iter_docx_blocks(document):
            if isinstance(block, Paragraph):
                text = block.text.strip()
                if text:
                    ordered.append(text)
            else:
                rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in block.rows]
                table_text = "\n".join(row for row in rows if row.strip())
                if table_text:
                    ordered.append(table_text)
        return [("document", "\n\n".join(ordered))]
    raise ValueError(f"Unsupported file type: {path}")


def _markdown_sections(text: str) -> list[tuple[str, str]]:
    """Split Markdown at ``<!-- page: N -->`` markers (written by ``convert-pdf``) into page sections."""
    from .pdf_markdown import PAGE_MARKER

    parts = PAGE_MARKER.split(text)
    if len(parts) == 1:
        return [("document", text)]
    sections = [(f"page {number}", body) for number, body in zip(parts[1::2], parts[2::2])]
    if parts[0].strip():  # front matter stays with the first page, where title lookups expect it
        sections[0] = (sections[0][0], parts[0] + sections[0][1])
    return sections


def _iter_docx_blocks(parent: DocumentObject | _Cell):
    """Yield DOCX paragraphs and tables in their source XML order."""
    parent_element = parent.element.body if isinstance(parent, DocumentObject) else parent._tc
    for child in parent_element.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, parent)
        elif child.tag.endswith("}tbl"):
            yield Table(child, parent)


def _blocks(text: str) -> list[str]:
    normalized_lines = [" ".join(line.split()) for line in text.splitlines()]
    blocks: list[str] = []
    current: list[str] = []
    for line in normalized_lines:
        heading = line.startswith("#") or (line.endswith(":") and len(line) < 160)
        if not line or heading:
            if current:
                blocks.append(" ".join(current))
                current = []
            if heading:
                blocks.append(line)
            continue
        current.append(line)
    if current:
        blocks.append(" ".join(current))
    return [block for block in blocks if block]


def _structured_chunks(text: str, chunk_chars: int, overlap_chars: int) -> list[tuple[str, str]]:
    if chunk_chars <= 0:
        raise ValueError("chunk size must be positive")
    if overlap_chars < 0:
        raise ValueError("chunk overlap cannot be negative")
    if overlap_chars >= chunk_chars:
        raise ValueError("chunk overlap must be smaller than chunk size")
    blocks = _blocks(text)
    if not blocks:
        return []
    expanded: list[str] = []
    for block in blocks:
        if len(block) <= chunk_chars:
            expanded.append(block)
            continue
        step = chunk_chars - overlap_chars
        expanded.extend(block[start : start + chunk_chars] for start in range(0, len(block), step))

    chunks: list[tuple[str, str]] = []
    start = 0
    while start < len(expanded):
        selected: list[str] = []
        used = 0
        end = start
        while end < len(expanded):
            addition = len(expanded[end]) + (2 if selected else 0)
            if selected and used + addition > chunk_chars:
                break
            selected.append(expanded[end])
            used += addition
            end += 1
        content = "\n\n".join(selected)
        # A short section can contain the only operative rule in a document.
        # chunk_chars is a bound/target, never an evidence deletion threshold.
        if content.strip():
            chunks.append((f"blocks {start + 1}-{end}", content))
        if end >= len(expanded):
            break
        overlap_start = end
        overlap_used = 0
        while overlap_start > start and overlap_used + len(expanded[overlap_start - 1]) <= overlap_chars:
            overlap_start -= 1
            overlap_used += len(expanded[overlap_start]) + 2
        start = overlap_start if overlap_start > start else end
    return chunks


def load_chunks(
    root: Path, chunk_chars: int, overlap_chars: int, progress_enabled: bool = False,
    extraction: dict | None = None,
) -> list[Chunk]:
    """Chunk every supported file under ``root``.

    ``extraction``, when supplied, receives an extraction summary: the number of files, the
    sections (PDF pages) that yielded no text, and the files that yielded none at all. Scanned PDFs
    show up there; they need OCR before Alloy can use them.
    """
    if chunk_chars <= 0:
        raise ValueError("chunk size must be positive")
    if overlap_chars < 0:
        raise ValueError("chunk overlap cannot be negative")
    if overlap_chars >= chunk_chars:
        raise ValueError("chunk overlap must be smaller than chunk size")
    files = corpus_files(root)
    if not files:
        raise ValueError(f"No supported documents found under {root}")
    chunks: list[Chunk] = []
    empty_sections: dict[str, list[str]] = {}
    files_without_text: list[str] = []
    for path in track(files, enabled=progress_enabled, description="Reading documents", total=len(files)):
        relative = str(path.relative_to(root))
        chunk_number = 0
        for section_location, raw_text in _read_sections(path):
            if not raw_text.strip():
                import logging
                logging.getLogger(__name__).warning("document_section_empty file=%s section=%s", relative, section_location)
                empty_sections.setdefault(relative, []).append(section_location)
            for block_location, part in _structured_chunks(raw_text, chunk_chars, overlap_chars):
                chunk_number += 1
                location = f"{section_location}, {block_location}"
                chunks.append(Chunk(f"{relative}#chunk-{chunk_number}", relative, location, part))
        if not chunk_number:
            files_without_text.append(relative)
    if extraction is not None:
        extraction.update(
            files=len(files), files_without_text=files_without_text, empty_sections=empty_sections,
        )
    if not chunks:
        raise ValueError("Documents contained no extractable text")
    return chunks
